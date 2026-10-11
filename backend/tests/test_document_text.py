import io

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from services.document_text import csv_to_text, docx_to_text, pptx_to_text, to_text, xlsx_to_text


def test_pptx_to_text_extracts_titles_and_body():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Test Deck"
    slide.placeholders[1].text = "Some bullet content"
    buf = io.BytesIO()
    prs.save(buf)

    text = pptx_to_text(buf.getvalue())
    assert "Test Deck" in text
    assert "Some bullet content" in text


def test_docx_to_text_extracts_paragraphs():
    doc = DocxDocument()
    doc.add_paragraph("One-pager summary text.")
    buf = io.BytesIO()
    doc.save(buf)

    assert "One-pager summary text." in docx_to_text(buf.getvalue())


def test_xlsx_to_text_extracts_rows():
    wb = Workbook()
    ws = wb.active
    ws.append(["Month", "Revenue"])
    ws.append(["Jan", 1000])
    buf = io.BytesIO()
    wb.save(buf)

    text = xlsx_to_text(buf.getvalue())
    assert "Month | Revenue" in text
    assert "Jan | 1000" in text


def test_csv_to_text():
    text = csv_to_text(b"name,sector\nAcme,FinTech\n")
    assert text == "name | sector\nAcme | FinTech"


def test_to_text_dispatches_by_kind():
    assert to_text("csv", b"a,b\n1,2\n") == "a | b\n1 | 2"


def test_to_text_rejects_pdf():
    with pytest.raises(ValueError):
        to_text("pdf", b"%PDF-1.4")


def test_xlsx_to_text_raises_on_corrupt_content():
    with pytest.raises(Exception):
        xlsx_to_text(b"this is not a zip file at all")
