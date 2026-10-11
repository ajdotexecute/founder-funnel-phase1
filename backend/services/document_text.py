"""
Converts non-PDF documents to plain text server-side before sending to Claude.
PDFs are sent natively (as a document content block) — see services/claude_client.py.
"""
import csv
import io

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation


def pptx_to_text(contents: bytes) -> str:
    prs = Presentation(io.BytesIO(contents))
    lines = []
    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def docx_to_text(contents: bytes) -> str:
    doc = DocxDocument(io.BytesIO(contents))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def xlsx_to_text(contents: bytes) -> str:
    wb = load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                lines.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


def csv_to_text(contents: bytes) -> str:
    text = contents.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)


CONVERTERS = {
    "pptx": pptx_to_text,
    "docx": docx_to_text,
    "xlsx": xlsx_to_text,
    "csv": csv_to_text,
}


def to_text(kind: str, contents: bytes) -> str:
    """Raises ValueError for 'pdf' — PDFs are sent natively, never converted."""
    if kind not in CONVERTERS:
        raise ValueError(f"No text converter for kind={kind!r}")
    return CONVERTERS[kind](contents)
